import csv
import io

import openpyxl

from app.models import ArtifactType
from app.parsers.base import BaseParser, ParseResult

XLSX_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
}
CSV_MIME_TYPES = {"text/csv", "application/csv"}


class SpreadsheetParser(BaseParser):
    def can_handle(self, mime_type: str) -> bool:
        return mime_type in XLSX_MIME_TYPES or mime_type in CSV_MIME_TYPES

    async def parse(self, file_path: str, mime_type: str) -> ParseResult:
        warnings: list[str] = []
        parts: list[str] = []
        sheet_count = 0
        row_count = 0

        if mime_type in XLSX_MIME_TYPES:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            sheet_count = len(wb.sheetnames)
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                parts.append(f"=== Sheet: {sheet_name} ===")
                for row in ws.iter_rows(values_only=True):
                    non_empty = [str(cell) for cell in row if cell is not None]
                    if non_empty:
                        parts.append("\t".join(non_empty))
                        row_count += 1
            wb.close()
        else:
            # CSV fallback
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for row in reader:
                    if any(cell.strip() for cell in row):
                        parts.append("\t".join(row))
                        row_count += 1
            sheet_count = 1

        full_text = "\n".join(parts).strip()

        return ParseResult(
            text=full_text or None,
            metadata={"sheet_count": sheet_count, "row_count": row_count},
            warnings=warnings,
            artifact_type=ArtifactType.SPREADSHEET_TEXT.value,
        )
